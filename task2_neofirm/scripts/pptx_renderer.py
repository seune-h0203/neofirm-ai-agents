#!/usr/bin/env python3
"""
pptx_renderer.py — Build NeoFirm Investment Committee PPTX from slides_data.json.

Usage:
  python pptx_renderer.py task2_neofirm/workspace/presentation/slides_data.json
"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    print("ERROR: python-pptx not installed. Run: pip install python-pptx", file=sys.stderr)
    sys.exit(1)

# Colour palette — light consulting theme
BG_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)   # white background
BG_SOFT   = RGBColor(0xF8, 0xF8, 0xF8)  # table header / soft surface (near-white)
ACCENT    = RGBColor(0x5B, 0x6C, 0xFF)  # blue #5B6CFF
FG_DARK   = RGBColor(0x17, 0x17, 0x17)  # main text
FG_MID    = RGBColor(0x69, 0x73, 0x86)  # secondary text
DIVIDER   = RGBColor(0xE0, 0xE2, 0xE8)  # rule / border
GOLD      = RGBColor(0xB4, 0x53, 0x09)  # amber (HOLD)
GREEN     = RGBColor(0x05, 0x96, 0x69)  # green (INVEST)
RED       = RGBColor(0x99, 0x1B, 0x1B)  # red (PASS)
SLIDE_W   = Inches(13.33)
SLIDE_H   = Inches(7.5)

DECISION_COLOR = {
    "INVEST": GREEN,
    "HOLD":   GOLD,
    "PASS":   RED,
}


def _bg(slide, color=BG_WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _box(slide, text, left, top, w, h,
         size=18, bold=False, color=FG_DARK,
         align=PP_ALIGN.LEFT, wrap=True):
    txb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = txb.text_frame
    tf.word_wrap = wrap
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _rule(slide, top, color=DIVIDER):
    shape = slide.shapes.add_shape(1, Inches(0.5), Inches(top), Inches(12.33), Inches(0.02))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _title_slide(prs, title, subtitle, sector, decision):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _box(slide, "NeoFirm", 0.5, 0.9, 12.33, 1.2,
         size=48, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    _box(slide, "AI 투자 위원회", 0.5, 2.2, 12.33, 0.7,
         size=22, color=FG_MID, align=PP_ALIGN.CENTER)
    _rule(slide, 3.1)
    _box(slide, title, 0.5, 3.3, 12.33, 0.8,
         size=22, bold=True, color=FG_DARK, align=PP_ALIGN.CENTER)
    _box(slide, f"분야: {sector}", 0.5, 4.2, 12.33, 0.6,
         size=16, color=FG_MID, align=PP_ALIGN.CENTER)
    if decision:
        dc = DECISION_COLOR.get(decision.upper(), FG_DARK)
        _box(slide, decision, 0.5, 5.0, 12.33, 0.9,
             size=32, bold=True, color=dc, align=PP_ALIGN.CENTER)
    _box(slide, datetime.now().strftime("%Y-%m-%d"), 0.5, 6.5, 12.33, 0.5,
         size=13, color=FG_MID, align=PP_ALIGN.CENTER)


def _content_slide(prs, title, points, body="", badge=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    if badge:
        bc = DECISION_COLOR.get(badge.upper(), ACCENT)
        _box(slide, badge, 10.5, 0.25, 2.3, 0.55,
             size=15, bold=True, color=bc, align=PP_ALIGN.CENTER)
    _box(slide, title, 0.5, 0.3, 9.8, 0.75,
         size=26, bold=True, color=ACCENT)
    _rule(slide, 1.15)
    y = 1.35
    for pt in points[:8]:
        _box(slide, f"• {pt}", 0.7, y, 11.7, 0.6, size=16, color=FG_DARK)
        y += 0.60
    if body:
        _box(slide, body[:400], 0.7, max(y + 0.1, 5.4), 11.7, 1.3,
             size=13, color=FG_MID)


def _scorecard_slide(prs, scores_by_reviewer, decisions):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    _box(slide, "리뷰어 평가표 (Reviewer Scorecard)", 0.5, 0.3, 12, 0.75,
         size=26, bold=True, color=ACCENT)
    _rule(slide, 1.15)
    reviewers = list(scores_by_reviewer.keys())
    col_w = 3.5
    # Header row background
    from pptx.util import Inches as _In
    hdr_bg = slide.shapes.add_shape(1, _In(0.5), _In(1.22), _In(12.33), _In(0.48))
    hdr_bg.fill.solid(); hdr_bg.fill.fore_color.rgb = BG_SOFT
    hdr_bg.line.fill.background()
    # Headers
    _box(slide, "평가 기준", 0.5, 1.25, 3.0, 0.45, size=13, bold=True, color=FG_MID)
    for i, rv in enumerate(reviewers):
        label = {"market": "시장성", "technology": "기술성", "business": "사업성"}.get(rv.lower(), rv.upper())
        _box(slide, label, 3.6 + i * col_w, 1.25, col_w - 0.1, 0.45,
             size=13, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    _rule(slide, 1.75)
    # Scores
    criteria = sorted({c for rv in scores_by_reviewer.values() for c in rv})
    y = 1.85
    for idx, crit in enumerate(criteria[:8]):
        _box(slide, crit.replace("_", " ").title(), 0.5, y, 3.0, 0.5, size=13, color=FG_DARK)
        for i, rv in enumerate(reviewers):
            val = scores_by_reviewer[rv].get(crit, "—")
            _box(slide, str(val), 3.6 + i * col_w, y, col_w - 0.1, 0.5,
                 size=13, color=FG_DARK, align=PP_ALIGN.CENTER)
        _rule(slide, y + 0.48)
        y += 0.50
    # Decision row
    _rule(slide, y + 0.08)
    dec_bg = slide.shapes.add_shape(1, _In(0.5), _In(y + 0.12), _In(12.33), _In(0.55))
    dec_bg.fill.solid(); dec_bg.fill.fore_color.rgb = BG_SOFT
    dec_bg.line.fill.background()
    _box(slide, "최종 판단", 0.5, y + 0.18, 3.0, 0.55, size=14, bold=True, color=FG_MID)
    for i, rv in enumerate(reviewers):
        dec = decisions.get(rv, "—").upper()
        dc = DECISION_COLOR.get(dec, FG_DARK)
        _box(slide, dec, 3.6 + i * col_w, y + 0.18, col_w - 0.1, 0.55,
             size=14, bold=True, color=dc, align=PP_ALIGN.CENTER)


def render(data: dict, out_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    title = data.get("title", "NeoFirm Investment Committee")
    sector = data.get("sector", "")
    decision = data.get("decision", "")

    # Slide 1: Cover
    _title_slide(prs, title, "", sector, decision)

    # Content slides
    for s in data.get("slides", []):
        badge = s.get("badge", "")
        _content_slide(prs, s["title"], s.get("key_points", []),
                       s.get("body", ""), badge)

    # Scorecard (if present)
    scores = data.get("scores_by_reviewer", {})
    decisions = data.get("decisions_by_reviewer", {})
    if scores:
        _scorecard_slide(prs, scores, decisions)

    prs.save(str(out_path))
    return out_path


def main():
    if len(sys.argv) < 2:
        print("Usage: pptx_renderer.py <slides_data.json>", file=sys.stderr)
        sys.exit(1)

    json_path = Path(sys.argv[1])
    if not json_path.exists():
        print(f"ERROR: {json_path} not found", file=sys.stderr)
        sys.exit(1)

    data = json.loads(json_path.read_text(encoding="utf-8-sig"))
    out_dir = json_path.parent
    safe = "".join(c if c.isalnum() or c in "-_ " else "_" for c in data.get("sector", "neofirm"))[:30]
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    out_path = out_dir / f"final_{safe}_{ts}.pptx"

    result = render(data, out_path)
    print(f"[OK] PPTX saved: {result}")

    # Copy to project output/
    import shutil
    out_project = Path(__file__).parent.parent.parent / "output" / out_path.name
    out_project.parent.mkdir(exist_ok=True)
    shutil.copy2(result, out_project)
    print(f"[OK] Copied to: {out_project}")


if __name__ == "__main__":
    main()
